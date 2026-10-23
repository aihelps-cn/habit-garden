#!/usr/bin/env python3
"""小小成长家 - Word 产品介绍生成器"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

doc = Document()

# ============ 全局样式 ============
style = doc.styles['Normal']
font = style.font
font.name = 'Noto Sans SC'
font.size = Pt(11)
style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Noto Sans SC')

GREEN = RGBColor(0x47, 0x6B, 0x50)
DARK = RGBColor(0x31, 0x3E, 0x35)
MUTED = RGBColor(0x91, 0x98, 0x8E)
ORANGE = RGBColor(0xDB, 0xA0, 0x4D)

def set_cell_shading(cell, color_hex):
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), color_hex)
    shading.set(qn('w:val'), 'clear')
    cell._tc.get_or_add_tcPr().append(shading)

def add_heading_styled(text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = GREEN
    return h

def add_bullet(text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    if bold_prefix:
        run = p.add_run(bold_prefix)
        run.bold = True
        run.font.color.rgb = GREEN
        p.add_run(text)
    else:
        p.add_run(text)
    return p

def add_feature_card(icon, title, desc):
    """模拟卡片式的段落"""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.left_indent = Cm(0.3)
    
    run_icon = p.add_run(f'{icon}  ')
    run_icon.font.size = Pt(14)
    
    run_title = p.add_run(title)
    run_title.bold = True
    run_title.font.size = Pt(12)
    run_title.font.color.rgb = GREEN
    
    p.add_run('\n')
    run_desc = p.add_run(desc)
    run_desc.font.size = Pt(10.5)
    run_desc.font.color.rgb = DARK
    return p

# ============ 封面 ============
doc.add_paragraph()
doc.add_paragraph()

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run('小小成长家')
run.font.size = Pt(44)
run.font.bold = True
run.font.color.rgb = GREEN

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run('LITTLE HABIT GARDEN')
run.font.size = Pt(14)
run.font.color.rgb = MUTED
run.font.italic = True
run.font.name = 'Noto Sans SC'

tagline = doc.add_paragraph()
tagline.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = tagline.add_run('用陪伴浇灌，让好习惯慢慢发芽')
run.font.size = Pt(16)
run.font.color.rgb = DARK

doc.add_paragraph()

# 分隔线
line = doc.add_paragraph()
line.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = line.add_run('🌱  ·  📖  ·  🌳  ·  🎁  ·  🔥')
run.font.size = Pt(18)

doc.add_page_break()

# ============ 目录页 ============
add_heading_styled('目录', level=1)
toc_items = [
    '一、产品理念',
    '二、核心功能',
    '三、产品特色',
    '四、典型使用场景',
    '五、技术架构',
    '六、版本信息',
]
for item in toc_items:
    p = doc.add_paragraph(item)
    p.paragraph_format.space_after = Pt(4)

doc.add_page_break()

# ============ 一、产品理念 ============
add_heading_styled('一、产品理念', level=1)

p = doc.add_paragraph()
p.add_run('我们相信，')
run = p.add_run('每一个好习惯，都是一颗种子；每一次认真的努力，都值得被看见。')
run.bold = True
p.add_run('\n\n小小成长家，是一款帮助孩子养成好习惯、让家长温柔陪伴的家庭工具。')
p.add_run('\n它不是任务打卡器，而是一片成长的花园——\n\n')
run_child = p.add_run('孩子端')
run_child.bold = True
p.add_run('记录自己每一天的认真，从"我做到了"里收获自信；\n')
run_parent = p.add_run('家长端')
run_parent.bold = True
p.add_run('看见孩子的每一份努力，用鼓励和认可浇灌成长。')

doc.add_paragraph()
add_bullet('不惩罚、不攀比、不焦虑', '核心理念：')
add_bullet('用正向激励代替催促', '积分机制：')
add_bullet('让成长有迹可循', '成长日历：')
add_bullet('可绑定多位家长共同陪伴', '家庭模式：')

doc.add_page_break()

# ============ 二、核心功能 ============
add_heading_styled('二、核心功能', level=1)

# 2.1 孩子端
add_heading_styled('2.1 孩子端', level=2)

features_child = [
    ('📖', '每日任务记录', '孩子可以查看今天要完成的任务（阅读、运动、家务…），完成后一键记录，并写下想对家长说的话。'),
    ('✅', '诚实打卡机制', '记录后不急于计分，先提交给家长审核。诚实地"没做完"也不是失败，调整节奏继续努力就好。'),
    ('🌿', '自主提交与自动提交', '习惯养成早期，系统每天 22:00 自动提交；孩子熟悉后也可主动提交。23:30 未审核自动通过。'),
    ('🎁', '心愿兑换', '积分可以兑换真实的小奖励——冰淇淋、看电视、现金……每一份努力都能换成真实的快乐。'),
    ('📅', '成长日历', 'GitHub 风格的热力图，一眼看见哪些天做得好、哪些天还可以加油。连续打卡天数一目了然。'),
]

for icon, title, desc in features_child:
    add_feature_card(icon, title, desc)

doc.add_paragraph()

# 2.2 家长端
add_heading_styled('2.2 家长端', level=2)

features_parent = [
    ('🌱', '任务管理', '预设了 30+ 常见任务模板（阅读、家务、运动、习惯），一键添加；也可自定义属于孩子的专属任务。'),
    ('👀', '温柔审核', '浏览孩子的打卡记录，一键"通过并加分"或"给一点提醒"。系统还会把孩子想说的话展示给你。'),
    ('⭐', '积分调节', '可以手动加分奖励特殊表现，也可以温柔减分提醒。每一笔都有记录，透明可追溯。'),
    ('🎁', '奖励兑现', '孩子兑换心愿后，家长确认已兑现或已支付，积分系统保持准确。现金兑换 1:1，家长自行支付。'),
    ('👨‍👩‍👧', '多家长绑定', '一个孩子可以绑定多位家长（父母、祖父母），任意一方审核即可。也可同时陪伴多个孩子。'),
    ('📊', '数据看板', '今日入账积分、等待审核数量、累计通过次数、成长日历热力图——让陪伴有据可依。'),
]

for icon, title, desc in features_parent:
    add_feature_card(icon, title, desc)

doc.add_page_break()

# ============ 三、产品特色 ============
add_heading_styled('三、产品特色', level=1)

highlights = [
    ('设计理念', '不是"你必须做完这些"，而是"我们一起来做这些"。温和克制的配色，像翻阅一本自然手册。'),
    ('首次使用', '孩子注册 → 生成绑定码 → 家长输入绑定码，三步完成家庭连接。'),
    ('预设 vs 自定义', '30+ 精心撰写的预设任务和奖励，点击直接添加，省去重复输入；也支持完全自定义。'),
    ('自动更新', '安卓 App 内置自动更新机制——有新版本时弹出提示，一键下载安装，始终保持最新。'),
    ('隐私与安全', '服务端存储，支持持久化登录；请求校验、限流保护，数据安全有保障。'),
    ('跨平台', 'Web 网页端（桌面/移动浏览器）+ Android App（Capacitor 混合框架），同一账号多端同步。'),
]

for category, desc in highlights:
    p = doc.add_paragraph()
    run_label = p.add_run(f'【{category}】')
    run_label.bold = True
    run_label.font.color.rgb = ORANGE
    p.add_run(f'  {desc}')

doc.add_page_break()

# ============ 四、典型使用场景 ============
add_heading_styled('四、典型使用场景', level=1)

scenarios = [
    {
        'title': '场景一：每天的小约定',
        'steps': [
            '家长打开 App，选择"阅读课外书 20 分钟"预设，一键添加',
            '又添加了"刷牙洗脸"、"早睡早起"等生活习惯',
            '孩子端同步显示今天的 3 个任务',
            '放学回家，孩子完成每项任务后记录并写下一句话',
            '晚上家长收到提醒，审核打卡，通过后自动加分',
        ]
    },
    {
        'title': '场景二：周末的小奖励',
        'steps': [
            '家长预设了"冰淇淋 20 积分"、"看电影 50 积分"等心愿',
            '一周下来孩子攒了 80 积分',
            '孩子在心愿商店选择"看电影"，一键兑换',
            '家长在兑换记录里看到，周末安排好电影',
            '看完后家长在 App 里点击"确认已兑现"',
        ]
    },
    {
        'title': '场景三：长期的习惯养成',
        'steps': [
            '成长日历热力图记录孩子的每一天',
            '家长发现最近打卡变少，温和询问原因',
            '一起调整任务难度或换一个更喜欢的奖励',
            '渐渐地，热力图上有了连续的绿色小方块',
            '3 个月后回看，那些坚持的日子清晰可见',
        ]
    },
    {
        'title': '场景四：多家长陪伴',
        'steps': [
            '爸爸和妈妈分别用各自账号绑定孩子',
            '爸爸每天审核学习类任务，妈妈负责生活习惯',
            '任意一方审核通过即可，积分只入账一次',
            '周末爷爷奶奶也能看到孩子的成长记录',
        ]
    },
]

for s in scenarios:
    p = doc.add_paragraph()
    run = p.add_run(s['title'])
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = GREEN
    
    for i, step in enumerate(s['steps'], 1):
        p = doc.add_paragraph(f'{i}. {step}', style='List Number')
        p.paragraph_format.space_after = Pt(2)

doc.add_page_break()

# ============ 五、技术架构 ============
add_heading_styled('五、技术架构', level=1)

# 技术栈表格
add_heading_styled('5.1 技术栈', level=2)

table = doc.add_table(rows=6, cols=3)
table.style = 'Light Grid Accent 1'
table.alignment = WD_TABLE_ALIGNMENT.CENTER

headers = ['层级', '技术选型', '说明']
for i, h in enumerate(headers):
    cell = table.rows[0].cells[i]
    cell.text = h
    for p in cell.paragraphs:
        for run in p.runs:
            run.bold = True
            run.font.color.rgb = GREEN
    set_cell_shading(cell, 'EAF0E4')

rows_data = [
    ['前端（Web）', '原生 JavaScript + CSS', '无框架依赖，轻量高性能'],
    ['移动端', 'Capacitor 8 + Android', '共享前端资源，一套代码多端运行'],
    ['后端', 'Node.js + Express', 'RESTful API，ES Module'],
    ['数据库', 'MySQL 8', '事务保证数据一致性'],
    ['部署', 'Docker Compose + Nginx', '一键部署 HTTPS 服务'],
]

for r, row_data in enumerate(rows_data, 1):
    for c, val in enumerate(row_data):
        cell = table.rows[r].cells[c]
        cell.text = val

doc.add_paragraph()

# 架构图（文字版）
add_heading_styled('5.2 架构概览', level=2)

arch_lines = [
    '┌──────────────┐     ┌──────────────┐     ┌──────────────┐',
    '│  安卓 App    │     │  Web 浏览器   │     │   其他设备    │',
    '└──────┬───────┘     └──────┬───────┘     └──────┬───────┘',
    '       │                     │                     │',
    '       └─────────────────────┼─────────────────────┘',
    '                             │ HTTPS / CORS',
    '                             ▼',
    '                    ┌─────────────────────┐',
    '                    │   Nginx (reverse)   │',
    '                    └──────────┬──────────┘',
    '                               │',
    '                               ▼',
    '                    ┌─────────────────────┐',
    '                    │  Node.js Express    │',
    '                    │  (Auth / Business)  │',
    '                    └──────────┬──────────┘',
    '                               │',
    '                               ▼',
    '                    ┌─────────────────────┐',
    '                    │     MySQL 8 数据库    │',
    '                    └─────────────────────┘',
]

for line in arch_lines:
    p = doc.add_paragraph()
    run = p.add_run(line)
    run.font.name = 'Menlo'
    run.font.size = Pt(8.5)

doc.add_paragraph()

add_heading_styled('5.3 安全设计', level=2)
add_bullet('所有 API 强制 HTTPS，CORS 白名单校验')
add_bullet('请求频率限制（登录接口 15 分钟 30 次、普通接口 1 分钟 300 次）')
add_bullet('密码 bcrypt 哈希存储，绑定码 SHA-256 哈希校验')
add_bullet('会话 Cookie HttpOnly + Secure + SameSite 严格')
add_bullet('业务操作通过数据库事务保证一致性')

doc.add_page_break()

# ============ 六、版本信息 ============
add_heading_styled('六、版本信息', level=1)

table = doc.add_table(rows=5, cols=3)
table.style = 'Light Grid Accent 1'
table.alignment = WD_TABLE_ALIGNMENT.CENTER

headers = ['版本', '发布日期', '主要更新']
for i, h in enumerate(headers):
    cell = table.rows[0].cells[i]
    cell.text = h
    for p in cell.paragraphs:
        for run in p.runs:
            run.bold = True
            run.font.color.rgb = GREEN
    set_cell_shading(cell, 'EAF0E4')

rows_data = [
    ['3.1.1', '2026-10-21', '优化自定义任务表单：图标选择器可折叠；新增"运动/活动"任务分类'],
    ['3.1.0', '2026-10-20', '移动端底部导航栏；预设任务一键添加；安卓应用内自动更新机制'],
    ['3.0.0', '2026-10-18', '首次发布：双端账号、任务打卡、家长审核、积分奖励、成长日历'],
    ['—', '持续迭代中', '每月根据用户反馈快速优化'],
]

for r, row_data in enumerate(rows_data, 1):
    for c, val in enumerate(row_data):
        cell = table.rows[r].cells[c]
        cell.text = val

doc.add_paragraph()

# 结尾
doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('「习惯养成不是一次冲刺，而是一次温柔的旅行。」')
run.font.size = Pt(12)
run.font.italic = True
run.font.color.rgb = MUTED

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('小小成长家 · 用陪伴浇灌，让好习惯慢慢发芽')
run.font.size = Pt(13)
run.font.bold = True
run.font.color.rgb = GREEN

# ============ 保存 ============
out_dir = os.path.dirname(os.path.abspath(__file__))
out_path = os.path.join(out_dir, '小小成长家_产品介绍.docx')
doc.save(out_path)
print(f'✅ 已生成：{out_path}')
print(f'   文件大小：{os.path.getsize(out_path) / 1024:.1f} KB')
